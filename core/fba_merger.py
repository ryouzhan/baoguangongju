import os
import re
import shutil
from datetime import datetime
import pandas as pd
import openpyxl
from core.config import load_columns_config, DEFAULT_CONFIG

class FbaCustomsMerger:
    def __init__(self):
        self.file_path = None
        self.df = None
        self.address_list = []
        self.address_mapping = {}
        self.fba_display_map = {}
        self.merge_groups = []
        self.config = load_columns_config().get("fba_merger", DEFAULT_CONFIG["fba_merger"])

    def load_file(self, file_path):
        try:
            if not os.path.exists(file_path):
                return False, "文件不存在"
            
            self.config = load_columns_config().get("fba_merger", DEFAULT_CONFIG["fba_merger"])
            addr_col = self.config.get("address_col", "物流中心编码")
            fba_col = self.config.get("fba_col", "FBA编号")
            prod_col = self.config.get("product_col", "中文品名")
            unit_col = self.config.get("unit_col", "单位")
            must_cols = self.config.get("must_cols", [addr_col, prod_col, unit_col])
            num_cols = self.config.get("num_cols", ['箱数','数量','金额USD','毛重KGS','净重KGS','体积','货值','单价'])

            df_raw = pd.read_excel(file_path, sheet_name="汇总", header=1)
            
            df_raw.columns = [str(c).strip() if pd.notna(c) else f"col_{i}" 
                              for i, c in enumerate(df_raw.columns)]
            
            for c in must_cols:
                if c not in df_raw.columns:
                    return False, f"表格格式错误：缺少必需列【{c}】"
            
            df_clean = df_raw.dropna(how='all').reset_index(drop=True)
            
            for col in num_cols:
                if col in df_clean.columns:
                    df_clean[col] = pd.to_numeric(
                        df_clean[col].astype(str).str.replace(',', '').str.strip(),
                        errors='coerce'
                    ).fillna(0)
            
            str_cols = [addr_col, prod_col, unit_col, fba_col]
            for col in str_cols:
                if col in df_clean.columns:
                    df_clean[col] = df_clean[col].astype(str).str.strip()

            df_clean[addr_col] = df_clean[addr_col].astype(str).apply(
                lambda x: re.sub(r'[\(（].*?[\)）]', '', x).strip()
            )
            
            df_clean = df_clean[df_clean[addr_col].notna()]
            df_clean = df_clean[df_clean[addr_col] != 'nan']
            df_clean = df_clean[df_clean[addr_col] != '']
            
            self.address_list = []
            self.address_mapping = {}
            self.fba_display_map = {}
            seen_display = set()
            
            for _, row in df_clean.iterrows():
                addr = row[addr_col]
                fba = row.get(fba_col, '').strip() if fba_col in row else ''
                
                if fba and fba != 'nan':
                    display_str = f"{addr}({fba})"
                else:
                    display_str = addr
                
                if display_str not in seen_display:
                    seen_display.add(display_str)
                    self.address_list.append(display_str)
                    self.address_mapping[display_str] = addr
                    if fba and fba != '':
                        self.fba_display_map[fba] = display_str
            
            self.df = df_clean
            self.file_path = file_path
            return True, {
                "row_count": len(df_clean),
                "address_count": len(self.address_list),
                "address_list": self.address_list,
                "fba_list": list(self.fba_display_map.keys())
            }
        
        except Exception as e:
            return False, f"加载失败：{str(e)}"

    def process_group(self, target_items):
        try:
            addr_col = self.config.get("address_col", "物流中心编码")
            fba_col = self.config.get("fba_col", "FBA编号")
            prod_col = self.config.get("product_col", "中文品名")
            unit_col = self.config.get("unit_col", "单位")
            out_addr_col = self.config.get("output_address_col", "配送地址")

            valid_addrs = [self.address_mapping[item] for item in target_items if item in self.address_mapping]
            
            if not valid_addrs:
                raise ValueError("未匹配到任何有效地址")
            
            filtered = self.df[self.df[addr_col].isin(valid_addrs)].copy()
            if filtered.empty:
                raise ValueError("未匹配到任何数据")
            
            full_fba_list = []
            for item in target_items:
                if '(' in item and ')' in item:
                    start_idx = item.rfind('(') + 1
                    end_idx = item.rfind(')')
                    if end_idx > start_idx:
                        fba_code = item[start_idx:end_idx].strip()
                        full_fba_list.append(fba_code)
            FULL_FBA = ' '.join(full_fba_list).strip()

            agg_dict = {
                '箱数': 'sum',
                '数量': 'sum',
                '毛重KGS': 'sum',
                '净重KGS': 'sum',
                '体积': 'sum',
                '货值': 'sum',
                '金额USD': 'sum',
                'HS编码': lambda x: ' '.join(pd.Series(x).dropna().astype(str).unique()),
                fba_col: lambda x: FULL_FBA,
                '中英文品名': lambda x: ' '.join(pd.Series(x).dropna().astype(str).unique()),
                '申报要素(品牌,材质,用途)': lambda x: ' '.join(pd.Series(x).dropna().astype(str).unique()),
                'CTNS': lambda x: ' '.join(pd.Series(x).dropna().astype(str).unique()),
                '币制': lambda x: ' '.join(pd.Series(x).dropna().astype(str).unique()),
                '境内货源地': lambda x: ' '.join(pd.Series(x).dropna().astype(str).unique()),
                '供应商': lambda x: ' '.join(pd.Series(x).dropna().astype(str).unique()),
                '合同号': lambda x: ' '.join(pd.Series(x).dropna().astype(str).unique()),
                '账号': lambda x: ' '.join(pd.Series(x).dropna().astype(str).unique())
            }

            agg_dict = {k: v for k, v in agg_dict.items() if k in filtered.columns}
            df_group = filtered.groupby([prod_col, unit_col], as_index=False).agg(agg_dict)
            
            float_cols = ['箱数','数量','毛重KGS','净重KGS','体积','货值','金额USD']
            for c in float_cols:
                if c in df_group.columns:
                    df_group[c] = df_group[c].round(2)
            
            if '单价' in df_group.columns:
                df_group['单价'] = ''
            
            df_group[out_addr_col] = ' '.join(target_items)
            return df_group
        
        except Exception as e:
            raise Exception(f"合并处理失败：{str(e)}")

    def export_merged_excel(self, groups, output_dir):
        if not self.file_path or not os.path.exists(self.file_path):
            raise ValueError("源文件不存在或尚未加载！")
        if not groups:
            raise ValueError("请至少提供一个待合并分组！")

        now = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_filename = f"FBA合并结果_{now}.xlsx"
        out_path = os.path.join(output_dir, out_filename)

        shutil.copy2(self.file_path, out_path)
        wb = openpyxl.load_workbook(out_path)
        if "汇总" not in wb.sheetnames:
            wb.close()
            raise ValueError("源文件中未找到工作表【汇总】")

        ws = wb["汇总"]
        total_merged_rows = 0

        for group in groups:
            df_merge = self.process_group(group)
            total_merged_rows += len(df_merge)
            for _, row in df_merge.iterrows():
                row_idx = ws.max_row + 1
                for col_idx, col_name in enumerate(self.df.columns, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=row.get(col_name, ""))

        wb.save(out_path)
        wb.close()

        return out_path, {
            "filename": out_filename,
            "groups_count": len(groups),
            "merged_rows_count": total_merged_rows
        }
