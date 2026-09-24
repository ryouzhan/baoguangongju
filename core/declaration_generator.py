import os
import re
from datetime import datetime
from copy import copy
import pandas as pd
import openpyxl
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, Border, Side, Alignment

FOOTER_KEYWORDS = [
    '合计', '总计', 'Total', 'SUM', 'Sum',
    '备注', 'Remark', 'Comments',
    '签字', 'Signature', 'Sign',
    '日期', 'Date',
    '申报单位', '填表人', 'Approved', 'Checked'
]

CLEAR_RULES = {
    1: {"报关单": (23, 34, 1, 30), "装箱单": (12, (2, 3, 4, 5), 1, 12),
        "发票": (15, (2, 3, 4, 5), 1, 25), "合同": (14, (2, 3, 4, 5), 1, 20)},
    2: {"报关单": (26, 34, 1, 30), "装箱单": (12, (3, 4, 5), 1, 12),
        "发票": (15, (3, 4, 5), 1, 25), "合同": (14, (3, 4, 5), 1, 20)},
    3: {"报关单": (29, 34, 1, 30), "装箱单": (12, (4, 5), 1, 12),
        "发票": (15, (4, 5), 1, 25), "合同": (14, (4, 5), 1, 20)},
    4: {"报关单": (32, 34, 1, 30), "装箱单": (12, (5,), 1, 12),
        "发票": (15, (5,), 1, 25), "合同": (14, (5,), 1, 20)},
    5: {"报关单": None, "装箱单": None, "发票": None, "合同": None}
}

THIN_BORDER = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
FONT_STYLE = Font(name='微软雅黑', size=10)

def normalize_header(value):
    if value is None: return ""
    return str(value).strip().replace(" ", "").replace("\n", "").replace("\t", "").replace("\r", "")

def find_real_header_row(file_path, sheet_name):
    try:
        df_temp = pd.read_excel(file_path, sheet_name=sheet_name, header=None, nrows=20, engine='openpyxl')
        for idx, row in df_temp.iterrows():
            row_vals = [normalize_header(val) for val in row if pd.notna(val)]
            if "FBA编号" in row_vals and "账号" in row_vals:
                return idx
    except Exception:
        pass
    return 0

def generate_unique_fba_code(index):
    now = datetime.now()
    year = now.strftime("%Y")
    month = now.strftime("%m")
    seq_num = str(index).zfill(4)
    return f"FBA-{year}-{month}-{seq_num}"

def get_top_left_cell(ws, cell):
    if isinstance(cell, MergedCell):
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                return ws.cell(row=merged_range.min_row, column=merged_range.min_col)
    return cell

def check_row_has_data(ws, row, col_indices):
    for col in col_indices:
        cell = ws.cell(row=row, column=col)
        real_cell = get_top_left_cell(ws, cell)
        if real_cell.value is not None and str(real_cell.value).strip() != "":
            return True, str(real_cell.value)
    return False, ""

def find_footer_boundary(ws, start_row, col_indices):
    max_row = ws.max_row
    empty_row_count = 0
    for r in range(start_row, max_row + 10):
        has_data, content = check_row_has_data(ws, r, col_indices)
        if has_data:
            content_norm = str(content).strip()
            for kw in FOOTER_KEYWORDS:
                if kw in content_norm: return r
            empty_row_count = 0
        else:
            empty_row_count += 1
            if empty_row_count >= 5: return r - 4
    return max_row + 1

def clear_data_range(ws, start_row, end_row, col_indices):
    if end_row < start_row: return
    for r in range(start_row, end_row):
        for c in col_indices:
            cell = ws.cell(row=r, column=c)
            if isinstance(cell, MergedCell): continue
            cell.value = None

def is_cell_merged(ws, row, col):
    for merged_range in ws.merged_cells.ranges:
        if (merged_range.min_row <= row <= merged_range.max_row) and \
           (merged_range.min_col <= col <= merged_range.max_col):
            return True
    return False

def clear_row_range(ws, start_row, end_row, start_col, end_col):
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            if is_cell_merged(ws, row, col):
                continue
            cell = ws.cell(row=row, column=col)
            cell.value = None
            cell.data_type = "n"
            if hasattr(cell, "border"):
                cell.border = copy(cell.border)

def clear_sequence_rows(ws, start_row, clear_seq, start_col, end_col):
    for seq in clear_seq:
        row_to_clear = start_row + (seq - 1)
        if row_to_clear > ws.max_row:
            continue
        for col in range(start_col, end_col + 1):
            if is_cell_merged(ws, row_to_clear, col):
                continue
            cell = ws.cell(row=row_to_clear, column=col)
            cell.value = None
            cell.data_type = "n"
            if hasattr(cell, "border"):
                cell.border = copy(cell.border)

def apply_template_cleanup(wb, actual_count, status_callback=None):
    """根据实际商品条数（1-5条），清理报关单、装箱单、发票、合同未用多余行"""
    if actual_count not in CLEAR_RULES:
        return
    
    rule_set = CLEAR_RULES[actual_count]
    bgd_rule = rule_set.get("报关单")
    if bgd_rule and "报关单" in wb.sheetnames:
        bgd_ws = wb["报关单"]
        r_start, r_end, c_start, c_end = bgd_rule
        clear_row_range(bgd_ws, r_start, r_end, c_start, c_end)

    for sheet in ("装箱单", "发票", "合同"):
        rule = rule_set.get(sheet)
        if rule and sheet in wb.sheetnames:
            r_start, seq_tuple, c_start, c_end = rule
            if seq_tuple:
                sheet_ws = wb[sheet]
                clear_sequence_rows(sheet_ws, r_start, seq_tuple, c_start, c_end)

def fill_and_clean_declaration(template_path, fba_data_df, fba_code, account_info, save_folder, progress_callback=None):
    """
    一步到位：填充模板、提取合同号、执行空白行清理，直接生成改名后的最终报关资料，无需产生或保留中间未清理文件。
    """
    wb = None
    try:
        wb = load_workbook(template_path, data_only=False, read_only=False)
        wb.calculation.calcMode = "manual"
        wb.calculation.calcOnSave = False
        if "输入表格" not in wb.sheetnames:
            raise ValueError("模板文件中未找到工作表【输入表格】")
        ws = wb["输入表格"]

        target_b12 = ws['B12']
        real_b12 = get_top_left_cell(ws, target_b12)
        real_b12.value = account_info

        field_mapping = {
            '合同号': '合同号', 'HS编码': 'HS编码', '中文品名': '中文品名', '中英文品名': '中英文品名',
            '申报要素(品牌,材质,用途)': '申报要素(品牌,材质,用途)', '箱数': '箱数', 'CTNS': 'CTNS',
            '数量': '数量', '单位': '单位', '单价': '单价', '金额USD': '金额USD',
            '毛重KGS': '毛重KGS', '净重KGS': '净重KGS', '体积': '体积', '境内货源地': '境内货源地',
            'FBA编号': 'FBA编号'
        }

        header_row = -1
        field_col = {}
        for r in range(1, 16):
            temp_field_col = {}
            match_count = 0
            for c in range(1, 51):
                cell_obj = ws.cell(row=r, column=c)
                real_cell = get_top_left_cell(ws, cell_obj)
                clean_val = normalize_header(real_cell.value)
                if clean_val in field_mapping:
                    temp_field_col[clean_val] = c
                    match_count += 1
            if match_count >= 5:
                header_row = r
                field_col = temp_field_col
                break

        if header_row == -1:
            raise ValueError("无法在模板的【输入表格】中识别表头字段！")

        data_start_row = header_row + 1
        mapped_col_indices = list(field_col.values())

        footer_boundary_row = find_footer_boundary(ws, data_start_row, mapped_col_indices)
        clear_data_range(ws, start_row=data_start_row, end_row=footer_boundary_row, col_indices=mapped_col_indices)

        df_clean = fba_data_df.fillna("")
        data_list = df_clean.values.tolist()

        col_order = df_clean.columns.tolist()
        normalized_col_order = [normalize_header(c) for c in col_order]
        field_data_idx = {}
        for tm_key, tm_name in field_mapping.items():
            try:
                field_data_idx[tm_key] = normalized_col_order.index(normalize_header(tm_name))
            except ValueError:
                pass

        total_rows = len(data_list)
        contract_no = ""

        for row_idx in range(total_rows):
            excel_row = data_start_row + row_idx
            row_data = data_list[row_idx]

            ws.row_dimensions[excel_row].height = 25

            for tm_field, dt_idx in field_data_idx.items():
                if tm_field in field_col:
                    col_idx = field_col[tm_field]
                    cell_value = row_data[dt_idx]

                    if tm_field == '合同号' and not contract_no and cell_value:
                        contract_no = str(cell_value).strip()

                    target_cell = ws.cell(row=excel_row, column=col_idx)

                    if isinstance(target_cell, MergedCell):
                        real_target = get_top_left_cell(ws, target_cell)
                        real_target.value = cell_value
                        real_target.border = THIN_BORDER
                        real_target.alignment = ALIGN_CENTER
                        real_target.font = FONT_STYLE
                    else:
                        target_cell.value = cell_value
                        target_cell.border = THIN_BORDER
                        target_cell.alignment = ALIGN_CENTER
                        target_cell.font = FONT_STYLE

            if progress_callback:
                progress_callback(int((row_idx + 1) / total_rows * 100))

        # 若未从输入表格填入的值提取到合同号，尝试从 grp 数据集中获取
        if not contract_no and '合同号' in fba_data_df.columns:
            contracts = fba_data_df['合同号'].dropna().astype(str).str.strip().tolist()
            if contracts and contracts[0]:
                contract_no = contracts[0]

        # 计算实际商品项数（最多5项）
        actual_items_count = min(max(total_rows, 1), 5)

        # 核心合并步骤：直接就地执行各联报关单、装箱单、发票、合同的空白模板行清理
        apply_template_cleanup(wb, actual_items_count)

        # 规范化命名：直接按 合同号 + 报关单 格式命名，不需要中间未改名文件
        safe_contract = re.sub(r'[/\\:*?"<>|]', '_', contract_no).strip() if contract_no else ""
        safe_fba = str(fba_code).replace("/", "_").replace("\\", "_").strip()
        safe_account = str(account_info).replace("/", "_").replace("\\", "_").strip()

        name_parts = []
        if safe_contract:
            name_parts.append(safe_contract)
        name_parts.append("报关单")
        if safe_account:
            name_parts.append(safe_account)
        if safe_fba:
            name_parts.append(safe_fba)

        final_filename = "_".join(name_parts) + ".xlsx"
        final_filepath = os.path.join(save_folder, final_filename)

        if os.path.exists(final_filepath):
            final_filename = "_".join(name_parts) + f"_{datetime.now().strftime('%H%M%S')}.xlsx"
            final_filepath = os.path.join(save_folder, final_filename)

        wb.calculation.calcMode = "auto"
        wb.save(final_filepath)
        return True, final_filepath, final_filename, actual_items_count
    except Exception as e:
        raise e
    finally:
        if wb:
            wb.close()

def process_all_fba_declarations(template_path, data_path, save_folder, progress_callback=None, status_callback=None):
    """
    遍历数据源中所有 FBA 分组，直接生成清理后并以合同号规范命名的报关单资料。
    """
    try:
        def status(m):
            if status_callback: status_callback(m)

        status("正在识别数据源表头...")
        header_idx = find_real_header_row(data_path, "汇总")
        status(f"在数据源第 {header_idx + 1} 行检测到表头...")

        df = pd.read_excel(data_path, sheet_name="汇总", header=header_idx, engine='openpyxl')
        df.columns = [normalize_header(c) for c in df.columns]

        required_fields = [
            '账号',
            '合同号',
            'HS编码', '中文品名', '中英文品名', '申报要素(品牌,材质,用途)',
            '箱数', 'CTNS', '数量', '单位', '单价', '金额USD',
            '毛重KGS', '净重KGS', '体积', '境内货源地', 'FBA编号'
        ]

        missing = [c for c in required_fields if c not in df.columns]
        if missing:
            raise ValueError(f"数据源表格缺少必需字段：{', '.join(missing)}")

        df = df[required_fields].dropna(how='all').reset_index(drop=True)
        if len(df) == 0:
            raise ValueError("数据源中无有效数据行！")

        fba_groups = df.groupby('FBA编号', dropna=False)
        if len(fba_groups) == 0:
            fba_groups = [("", df)]

        total = len(fba_groups)
        success_count = 0
        idx = 0
        generated_files = []

        status(f"发现 {total} 个 FBA 分组，开始批量套版并就地自动清理空白行与重命名...")

        for fba, grp in fba_groups:
            idx += 1
            final_fba = fba if (pd.notna(fba) and str(fba).strip()) else generate_unique_fba_code(idx)

            account_val = ""
            if '账号' in grp.columns and not grp['账号'].isnull().all():
                first_val = grp['账号'].iloc[0]
                account_val = str(first_val) if pd.notna(first_val) else ""

            grp = grp.copy()
            grp.loc[:, 'FBA编号'] = final_fba

            def cb(p):
                if progress_callback:
                    progress_callback(int((idx - 1 + p/100) / total * 100))

            ok, fpath, fname, items_cnt = fill_and_clean_declaration(
                template_path=template_path,
                fba_data_df=grp,
                fba_code=final_fba,
                account_info=account_val,
                save_folder=save_folder,
                progress_callback=cb
            )

            if ok:
                success_count += 1
                generated_files.append(fpath)
                status(f"[{idx}/{total}] ✅ 直接生成完成 (品项数:{items_cnt}项, 自动清除空白行): {fname}")

        return True, {
            "total_fba": total,
            "success_count": success_count,
            "files": generated_files
        }

    except Exception as e:
        return False, str(e)
