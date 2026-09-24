import os
import re
from copy import copy
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook

CLEAR_RULES = {
    1: {"报关单": (23, 34, 1, 30), "装箱单": (12, (2, 3, 4, 5), 1, 12),
        "发票": (15, (2, 3, 4, 5), 1, 25), "合同": (14, (2, 3, 4, 5), 1, 20)},
    2: {"报关单": (26, 34, 1, 30), "装箱单": (12, (3, 4, 5), 1, 12),
        "发票": (15, (3, 4, 5), 1, 25), "合同": (14, (3, 4, 5), 1, 20)},
    3: {"报关单": (29, 34, 1, 30), "装箱单": (12, (4, 5), 1, 12),
        "发票": (15, (4, 5), 1, 25), "合同": (14, (4, 5), 1, 20)},
    4: {"报关单": (32, 34, 1, 30), "装箱单": (12, (5,), 1, 12),
        "发票": (15, (5,), 1, 25), "合同": (14, (5,), 1, 20)},
    5: {"报关单": None, "装箱单": None, "发票": None, "合同": None}
}

def is_cell_merged(ws, row, col):
    for merged_range in ws.merged_cells.ranges:
        if (merged_range.min_row <= row <= merged_range.max_row) and \
           (merged_range.min_col <= col <= merged_range.max_col):
            return True
    return False

def count_actual_items(ws):
    count = 0
    for row in range(2, min(7, ws.max_row + 1)):
        cell = ws.cell(row=row, column=1)
        cell_value = str(cell.value).strip() if cell.value is not None else ""
        if cell_value and any(char.isdigit() for char in cell_value):
            count += 1
        else:
            if count > 0:
                break
    return min(count, 5)

def clear_row_range(ws, start_row, end_row, start_col, end_col):
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            if is_cell_merged(ws, row, col):
                continue
            cell = ws.cell(row=row, column=col)
            cell.value = None
            cell.data_type = "n"
            if hasattr(cell, "border"):
                cell.border = copy(cell.border)

def clear_sequence_rows(ws, start_row, clear_seq, start_col, end_col):
    for seq in clear_seq:
        row_to_clear = start_row + (seq - 1)
        if row_to_clear > ws.max_row:
            continue
        for col in range(start_col, end_col + 1):
            if is_cell_merged(ws, row_to_clear, col):
                continue
            cell = ws.cell(row=row_to_clear, column=col)
            cell.value = None
            cell.data_type = "n"
            if hasattr(cell, "border"):
                cell.border = copy(cell.border)

def clean_single_customs_file(file_path, output_dir, log_callback=None):
    def log(msg, level="info"):
        if log_callback:
            log_callback(msg, level)

    file_name_orig = os.path.basename(file_path)
    log(f"开始处理文件: {file_name_orig}")

    if file_path.endswith(".xls"):
        df_dict = pd.read_excel(file_path, sheet_name=None, engine="xlrd")
        temp_xlsx = os.path.join(output_dir, f"temp_{int(datetime.now().timestamp())}.xlsx")
        with pd.ExcelWriter(temp_xlsx, engine="openpyxl") as writer:
            for sheet_name, df in df_dict.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)
        wb = load_workbook(temp_xlsx, data_only=False)
        if os.path.exists(temp_xlsx):
            os.remove(temp_xlsx)
    else:
        wb = load_workbook(file_path, data_only=False)

    required_sheets = ["输入表格", "报关单", "装箱单", "发票", "合同"]
    missing_sheets = [s for s in required_sheets if s not in wb.sheetnames]
    if missing_sheets:
        wb.close()
        raise ValueError(f"缺少必要工作表：{','.join(missing_sheets)}")

    input_ws = wb["输入表格"]
    actual_count = count_actual_items(input_ws)
    if actual_count < 1:
        wb.close()
        raise ValueError("输入表格中未检测到有效商品数据（第2-6行第1列无有效编码/序号）")
    log(f"检测到有效商品项数: {actual_count} 项")

    contract_no = ""
    header_row = -1
    contract_col = -1

    for r in range(1, 10):
        for c in range(1, 30):
            cell = input_ws.cell(row=r, column=c)
            val = str(cell.value).strip() if cell.value is not None else ""
            if val == "合同号":
                header_row = r
                contract_col = c
                break
        if header_row != -1:
            break

    if contract_col != -1:
        for r in range(header_row + 1, header_row + 10):
            cell = input_ws.cell(row=r, column=contract_col)
            val = str(cell.value).strip() if cell.value is not None else ""
            if val:
                contract_no = val
                break

    if not contract_no:
        contract_no = os.path.splitext(file_name_orig)[0]

    contract_no = re.sub(r'[/\\:*?"<>|]', '_', contract_no).strip()

    if actual_count in CLEAR_RULES:
        rule_set = CLEAR_RULES[actual_count]
        bgd_rule = rule_set.get("报关单")
        if bgd_rule:
            bgd_ws = wb["报关单"]
            r_start, r_end, c_start, c_end = bgd_rule
            clear_row_range(bgd_ws, r_start, r_end, c_start, c_end)
            log(f"报关单：已清理第 {r_start}-{r_end} 行多余模板格")

        for sheet in ("装箱单", "发票", "合同"):
            rule = rule_set.get(sheet)
            if rule:
                r_start, seq_tuple, c_start, c_end = rule
                if seq_tuple:
                    sheet_ws = wb[sheet]
                    clear_sequence_rows(sheet_ws, r_start, seq_tuple, c_start, c_end)
                    log(f"{sheet}：已清理序号 {seq_tuple} 多余模板行")

    out_file_name = f"{contract_no}_{file_name_orig}"
    if not out_file_name.endswith(".xlsx"):
        out_file_name += ".xlsx"
    
    output_path = os.path.join(output_dir, out_file_name)
    wb.save(output_path)
    wb.close()

    log(f"✅ 处理完成并保存为: {out_file_name}")
    return output_path, out_file_name
